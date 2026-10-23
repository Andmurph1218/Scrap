"""Gera planilhas .xlsx por categoria folha alvo (Mercado Livre)."""

from __future__ import annotations

import json
import os
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from wago.config import (
    MAP_FILE_GENERAL,
    MAP_FILE_TARGET,
    ML_TEMPLATES_DIR,
    ensure_dirs,
)
from wago.ml.auth import api_call, refresh_ml_token

CAMPOS_GLOBAIS = [
    "SKU",
    "Título",
    "Preço [R$]",
    "Estoque",
    "Condição",
    "Descrição",
    "Fotos",
    "Marca",
    "Modelo",
]

FAMILIAS_ALVO = {
    "industria_comercio": {
        "keywords": [
            "industria",
            "indústria",
            "comercio",
            "comércio",
            "automacao",
            "automação",
            "componente",
            "equipamento",
            "maquina",
            "máquina",
            "maquinario",
            "maquinário",
        ],
        "required_roots": ["Indústria e Comércio"],
    },
    "ferramentas": {
        "keywords": [
            "ferramenta",
            "ferragens",
            "parafuso",
            "broca",
            "serra",
            "furadeira",
            "esmerilhadeira",
            "chave",
            "alicate",
        ],
        "required_roots": ["Ferramentas"],
    },
    "construcao": {
        "keywords": [
            "construcao",
            "construção",
            "obra",
            "banheiro",
            "eletrica",
            "elétrica",
            "iluminacao",
            "iluminação",
            "encanamento",
            "hidraulica",
            "hidráulica",
            "material",
            "tintas",
            "pisos",
            "revestimento",
            "portas",
            "janelas",
        ],
        "required_roots": ["Construção"],
    },
}


def normalize_text(text: str) -> str:
    if not text:
        return ""
    text = text.lower().strip()
    text = re.sub(r"[áàãâä]", "a", text)
    text = re.sub(r"[éèêë]", "e", text)
    text = re.sub(r"[íìîï]", "i", text)
    text = re.sub(r"[óòõôö]", "o", text)
    text = re.sub(r"[úùûü]", "u", text)
    text = re.sub(r"[ç]", "c", text)
    text = re.sub(r"[^a-z0-9 ]", "", text)
    return text


def sanitize_filename(name: str) -> str:
    name = re.sub(r"[^\w\s-]", "", name)
    name = re.sub(r"[-\s]+", "_", name).strip().strip("_")
    return name


def load_general_map() -> dict[str, Any] | None:
    if not MAP_FILE_GENERAL.exists():
        print(
            f"Arquivo {MAP_FILE_GENERAL} não encontrado.\n"
            "Execute antes: python -m wago ml-map"
        )
        return None
    return json.loads(MAP_FILE_GENERAL.read_text(encoding="utf-8"))


def load_target_mapping() -> dict[str, Any]:
    if MAP_FILE_TARGET.exists():
        try:
            return json.loads(MAP_FILE_TARGET.read_text(encoding="utf-8"))
        except Exception as exc:
            print(f"Erro ao carregar mapeamento existente: {exc}")
    return {}


def save_target_mapping(mapping: dict[str, Any]) -> None:
    ensure_dirs()
    MAP_FILE_TARGET.write_text(
        json.dumps(mapping, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def category_matches_families(
    category_info: dict[str, Any], families: dict
) -> list[str]:
    caminho = category_info.get("caminho", "")
    nome = category_info.get("nome", "")
    caminho_norm = normalize_text(caminho)
    nome_norm = normalize_text(nome)
    matched: list[str] = []

    for family_key, config in families.items():
        required_roots = config.get("required_roots", [])
        in_required_root = any(
            normalize_text(root) in caminho_norm for root in required_roots
        )
        if required_roots and not in_required_root:
            continue
        for keyword in config["keywords"]:
            if keyword in caminho_norm or keyword in nome_norm:
                matched.append(family_key)
                break
    return matched


def discover_manual_files() -> dict[str, list[str]]:
    manual_map: dict[str, list[str]] = {}
    if not ML_TEMPLATES_DIR.exists():
        return manual_map
    for root_dir, _dirs, files in os.walk(ML_TEMPLATES_DIR):
        for manual_file in files:
            if not manual_file.endswith(".xlsx"):
                continue
            manual_path = os.path.join(root_dir, manual_file)
            try:
                wb = load_workbook(manual_path, read_only=True, data_only=True)
                headers = [
                    str(cell.value).strip()
                    for cell in wb.active[1]
                    if cell.value
                ]
                wb.close()
            except Exception as exc:
                print(f"Erro ao ler planilha manual {manual_path}: {exc}")
                continue
            cat_id_match = re.match(r"(MLB\d+)", manual_file)
            if cat_id_match and headers:
                manual_map[cat_id_match.group(1)] = headers
    return manual_map


def build_folder_path(path_str: str) -> str:
    parts = [sanitize_filename(p) for p in path_str.split(" > ") if p.strip()]
    return os.path.join(str(ML_TEMPLATES_DIR), *parts)


def generate_excel(
    category_id: str,
    category_name: str,
    path_str: str,
    attributes: list[dict],
    manual_headers: list[str] | None = None,
    technical_specs: Any = None,
) -> dict[str, Any]:
    folder_path = build_folder_path(path_str)
    os.makedirs(folder_path, exist_ok=True)
    safe_name = sanitize_filename(category_name)
    filename = f"{category_id}_{safe_name}.xlsx"
    filepath = os.path.join(folder_path, filename)

    headers = list(CAMPOS_GLOBAIS)
    required_headers: list[str] = []

    if manual_headers:
        headers = manual_headers
    else:
        for attr in attributes:
            attr_name = attr.get("name", "")
            if attr_name and attr_name not in headers:
                headers.append(attr_name)
                if attr.get("tags", {}).get("required") or attr.get("required"):
                    required_headers.append(attr_name)
        if technical_specs:
            specs = technical_specs
            if isinstance(specs, dict):
                specs = specs.get("specs", specs.get("attributes", []))
            if isinstance(specs, list):
                for spec in specs:
                    spec_name = (
                        spec.get("name", "")
                        if isinstance(spec, dict)
                        else str(spec)
                    )
                    if spec_name and spec_name not in headers:
                        headers.append(spec_name)

    wb = Workbook()
    ws = wb.active
    ws.title = "Anuncios"
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            start_color="366092", end_color="366092", fill_type="solid"
        )
        if header in required_headers:
            cell.fill = PatternFill(
                start_color="C0504D", end_color="C0504D", fill_type="solid"
            )
    for col_idx, header in enumerate(headers, 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = min(
            max(len(header) + 4, 12), 50
        )
    wb.save(filepath)
    return {
        "filename": filename,
        "filepath": os.path.relpath(filepath),
        "headers": headers,
        "required_headers": required_headers,
        "total_columns": len(headers),
    }


def process_category(
    cat_id: str,
    cat_info: dict[str, Any],
    mapping: dict[str, Any],
    manual_map: dict[str, list[str]],
) -> None:
    cat_name = cat_info.get("nome", "")
    path_str = cat_info.get("caminho", "")
    print(f"\n[{len(mapping) + 1}] {cat_id} — {cat_name}")
    print(f"   Caminho: {path_str}")

    try:
        attributes = api_call(f"/categories/{cat_id}/attributes") or []
        if not isinstance(attributes, list):
            attributes = []
        technical_specs = api_call(f"/categories/{cat_id}/technical_specs")
        manual_headers = manual_map.get(cat_id)
        excel_info = generate_excel(
            cat_id,
            cat_name,
            path_str,
            attributes,
            manual_headers,
            technical_specs,
        )

        required_attrs = []
        optional_attrs = []
        for attr in attributes:
            attr_name = attr.get("name", "")
            if not attr_name:
                continue
            is_required = attr.get("tags", {}).get("required") or attr.get(
                "required"
            )
            (required_attrs if is_required else optional_attrs).append(attr_name)

        mapping[cat_id] = {
            "name": cat_name,
            "path": path_str,
            "filename": excel_info["filename"],
            "filepath": excel_info["filepath"],
            "total_columns": excel_info["total_columns"],
            "headers": excel_info["headers"],
            "required_attributes": required_attrs,
            "optional_attributes": optional_attrs,
            "global_fields": CAMPOS_GLOBAIS,
            "source": "manual" if manual_headers else "api_attributes",
        }
        save_target_mapping(mapping)
        print(f"   Salvo: {excel_info['filepath']}")
    except Exception as exc:
        print(f"   Erro ao processar {cat_id}: {exc}")


def gerar_planilhas_alvo(*, limit: int | None = None) -> None:
    ensure_dirs()
    token = refresh_ml_token()
    if not token:
        raise SystemExit("Token ML inválido.")

    general_map = load_general_map()
    if not general_map:
        return

    print(f"Mapa geral: {len(general_map)} categorias")
    mapping = load_target_mapping()
    manual_map = discover_manual_files()
    print(f"Planilhas manuais: {len(manual_map)}")

    target_leafs = []
    for cat_id, cat_info in general_map.items():
        if not cat_info.get("eh_folha"):
            continue
        matched = category_matches_families(cat_info, FAMILIAS_ALVO)
        if matched:
            target_leafs.append((cat_id, cat_info, matched))

    print(f"Folhas alvo: {len(target_leafs)}")
    processed = 0
    for cat_id, cat_info, _matched in target_leafs:
        if limit is not None and processed >= limit:
            break
        if cat_id in mapping:
            continue
        folder_path = build_folder_path(cat_info.get("caminho", ""))
        safe_name = sanitize_filename(cat_info.get("nome", ""))
        expected = os.path.join(folder_path, f"{cat_id}_{safe_name}.xlsx")
        if os.path.exists(expected):
            continue
        process_category(cat_id, cat_info, mapping, manual_map)
        processed += 1

    print(f"\nCategorias no mapeamento: {len(mapping)}")
    print(f"Mapeamento: {MAP_FILE_TARGET}")
