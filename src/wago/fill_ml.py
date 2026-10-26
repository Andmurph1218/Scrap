"""Etapa 5: preenche o Anunciar-ML.xlsx oficial com dados do scrape."""

from __future__ import annotations

import json
import re
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from lxml import etree
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from wago.config import OUT, ROOT, ensure_dirs
from wago.export import selecionar_para_exportacao


def _set_cell(ws, row: int, col: int, value: Any) -> None:
    cell = ws.cell(row, col)
    if isinstance(cell, MergedCell):
        return
    cell.value = value

TEMPLATE = ROOT / "templates" / "Anunciar-ML.xlsx"

# Headers canónicos (prefixo) → chave lógica
HEADER_ALIASES: list[tuple[str, str]] = [
    ("título:", "titulo"),
    ("titulo:", "titulo"),
    ("condição", "condicao"),
    ("condicao", "condicao"),
    ("código universal", "gtin"),
    ("codigo universal", "gtin"),
    ("varia por: nome comercial da cor", "varia_cor"),
    ("varia por: voltagem", "varia_voltagem"),
    ("varia por: gama de voltagem", "varia_gama_voltagem"),
    ("varia por: unidade de gama de voltagem", "varia_unidade_voltagem"),
    ("fotos", "fotos"),
    ("sku", "sku"),
    ("estoque", "estoque"),
    ("preço [r$]", "preco"),
    ("preco [r$]", "preco"),
    ("formato de venda", "formato_venda"),
    ("quantidade de disjuntores", "qtd_unidades"),
    ("quantidade de lâmpadas", "qtd_unidades"),
    ("quantidade de lampadas", "qtd_unidades"),
    ("quantidade de botões", "qtd_unidades"),
    ("quantidade de botoes", "qtd_unidades"),
    ("quantidade de tomadas de corriente", "qtd_unidades"),
    ("quantidade de conectores", "qtd_unidades"),
    ("quantidade de fusíveis", "qtd_unidades"),
    ("quantidade de fusiveis", "qtd_unidades"),
    ("quantidade de fichas elétricas", "qtd_unidades"),
    ("quantidade de fichas eletricas", "qtd_unidades"),
    ("descrição", "descricao"),
    ("descricao", "descricao"),
    ("largura (cm)", "largura"),
    ("altura (cm)", "altura"),
    ("profundidade (cm)", "profundidade"),
    ("peso físico", "peso"),
    ("peso fisico", "peso"),
    ("tipo de anúncio", "tipo_anuncio"),
    ("tipo de anuncio", "tipo_anuncio"),
    ("forma de envio", "forma_envio"),
    ("custo de envio", "custo_envio"),
    ("retirar pessoalmente", "retirar"),
    ("tipo de garantia", "tipo_garantia"),
    ("tempo de garantia", "tempo_garantia"),
    ("unidade de tempo de garantia", "unidade_garantia"),
    ("marca", "marca"),
    ("modelo alfanumérico", "modelo_alfanumerico"),
    ("modelo alfanumerico", "modelo_alfanumerico"),
    ("fabricante", "fabricante"),
    ("modelo", "modelo"),
    ("tipo de produto", "tipo_produto"),
    ("tipo de conector", "tipo_conector"),
    ("gênero do conector", "genero_conector"),
    ("genero do conector", "genero_conector"),
    ("material de revestimento", "material_revestimento"),
    ("tipo de tomada elétrica", "tipo_tomada"),
    ("tipo de tomada eletrica", "tipo_tomada"),
    ("tipo de conexão", "tipo_conexao_ficha"),
    ("tipo de conexao", "tipo_conexao_ficha"),
    ("quantidade de caracteres", "qtd_chars"),
]

ABA_AUTOMATICA = "Automático"

_MAPA_CORES_ML: list[tuple[str, str]] = [
    ("laranja claro", "Laranja-claro"),
    ("laranja escuro", "Laranja-escuro"),
    ("laranja", "Laranja"),
    ("cinza escuro", "Cinza-escuro"),
    ("cinza", "Cinza"),
    ("azul claro", "Azul-claro"),
    ("azul escuro", "Azul-escuro"),
    ("azul", "Azul"),
    ("verde claro", "Verde-claro"),
    ("verde escuro", "Verde-escuro"),
    ("verde-amarelo", "Verde"),
    ("verde", "Verde"),
    ("amarelo", "Amarelo"),
    ("vermelho", "Vermelho"),
    ("preto", "Preto"),
    ("branco", "Branco"),
    ("marrom claro", "Marrom-claro"),
    ("marrom escuro", "Marrom-escuro"),
    ("marrom", "Marrom"),
    ("violeta", "Violeta"),
    ("rosa", "Rosa"),
    ("bege", "Bege"),
    ("prateado", "Prateado"),
    ("prata", "Prateado"),
    ("dourado", "Dourado"),
]


def inferir_aba_ml(row: dict[str, Any], abas_disponiveis: list[str] | None = None) -> str:
    """Escolhe a aba do template ML correspondente à categoria/título do produto."""
    categorias = _norm(row.get("Categorias", ""))
    nome = _norm(row.get("Nome", ""))
    descricao = _norm(str(row.get("Descrição", ""))[:800])
    texto_principal = f"{categorias} {nome}"
    texto_completo = f"{texto_principal} {descricao}"

    if any(
        termo in texto_principal
        for termo in (
            "ferramenta de operacao",
            "ferramenta de operação",
            "ferramentas",
            "marcacao",
            "marcação",
            "cartao de identificacao",
            "cartão de identificação",
            "cartao de identficacao",
            "cartão de identficação",
            "protecao para os dedos",
            "proteção para os dedos",
            "tampa final",
            "placa final",
            "placa separadora",
            "sistemas de i/o",
            "entrada digital",
            "saida digital",
            "saída digital",
            "fonte de alimentacao",
            "fonte de alimentação",
            "controlador",
            "switch",
        )
    ):
        if abas_disponiveis and "Outros" not in abas_disponiveis:
            return abas_disponiveis[0]
        return "Outros"

    regras: list[tuple[str, tuple[str, ...]]] = [
        (
            "Interruptores Automáticos",
            (
                "disjuntor",
                "interruptor automatico",
                "interruptor automático",
                "interruptores automaticos",
                "interruptores automáticos",
                "circuit breaker",
                "protecao de circuito",
                "proteção de circuito",
            ),
        ),
        (
            "Lâmpadas",
            ("lampada", "lâmpada", "lampadas", "lâmpadas", "luminaria", "luminária"),
        ),
        ("Contatores", ("contator", "contatores")),
        (
            "Botões Industriais",
            (
                "botao",
                "botão",
                "botoes",
                "botões",
                "pulsador",
                "botoeira",
            ),
        ),
        ("Tomadas", ("tomada", "tomadas")),
        (
            "Relés",
            (
                "rele",
                "relé",
                "reles",
                "relés",
                "acoplador otico",
                "acoplador ótico",
                "acopladores oticos",
                "acopladores óticos",
                "optocoupler",
                "ssr",
            ),
        ),
        (
            "Fusíveis",
            (
                "fusivel",
                "fusível",
                "fusiveis",
                "fusíveis",
                "porta-fusivel",
                "porta-fusível",
                "borne fusivel",
                "borne fusível",
                "borne-fusivel",
                "borne-fusível",
            ),
        ),
        (
            "Medidores de Consumo Elétrico",
            (
                "medidor de energia",
                "medidores de energia",
                "medidor de consumo",
                "medidores de consumo",
                "energy meter",
            ),
        ),
        (
            "Interruptores Elétricos",
            (
                "interruptor eletrico",
                "interruptor elétrico",
                "interruptores eletricos",
                "interruptores elétricos",
                "chave seccionadora",
                "interruptor",
            ),
        ),
        (
            "Fichas Elétricas",
            (
                "ficha eletrica",
                "ficha elétrica",
                "fichas eletricas",
                "fichas elétricas",
                "plugue",
            ),
        ),
        (
            "Conectores",
            (
                "conector",
                "conectores",
                "borne",
                "bornes",
                "emenda",
                "jumper",
                "ponteira",
                "terminal",
                "regua de borne",
                "régua de borne",
                "regua macho",
                "régua macho",
                "regua femea",
                "régua fêmea",
                "cabo de conexao",
                "cabo de conexão",
                "cabo de interconexao",
                "cabo de interconexão",
                "cabo para sensor",
                "cabo de sistema",
                "derivador",
                "winsta",
                "topjob",
                "picoMAX",
                "mcs",
            ),
        ),
    ]

    escolhida = "Outros"
    for aba, palavras in regras:
        if any(p.lower() in texto_principal for p in palavras):
            escolhida = aba
            break
    else:
        for aba, palavras in regras:
            if any(p.lower() in texto_completo for p in palavras):
                escolhida = aba
                break

    if abas_disponiveis and escolhida not in abas_disponiveis:
        return "Outros" if "Outros" in abas_disponiveis else abas_disponiveis[0]
    return escolhida


def _extrair_cor_ml(row: dict[str, Any]) -> str:
    texto = f"{row.get('Nome', '')}\n{row.get('Descrição', '')}"
    m = re.search(r"•\s*Cor\s*:\s*([^\n•;]+)", texto, re.I)
    alvo = _norm(m.group(1)) if m else _norm(texto)
    for termo, cor_ml in _MAPA_CORES_ML:
        if termo in alvo:
            return cor_ml
    for termo, cor_ml in _MAPA_CORES_ML:
        if termo in _norm(texto):
            return cor_ml
    return "Cinza"


def _extrair_voltagem_ml(row: dict[str, Any]) -> str:
    texto = f"{row.get('Nome', '')}\n{row.get('Descrição', '')}"
    for padrao, saida in (
        (r"\b127\s*/\s*220\s*v\b", "127/220V"),
        (r"\b220\s*v\b", "220V"),
        (r"\b127\s*v\b", "127V"),
        (r"\b380\s*v\b", "380V"),
        (r"\b480\s*v\b", "480V"),
        (r"\b600\s*v\b", "600V"),
        (r"\b24\s*v(?:dc|ac|\b)", "24V"),
        (r"\b12\s*v(?:dc|ac|\b)", "12V"),
        (r"\b48\s*v(?:dc|ac|\b)", "48V"),
        (r"\b5\s*v(?:dc|ac|\b)", "5V"),
    ):
        if re.search(padrao, texto, re.I):
            return saida
    return ""


def _extrair_genero_conector(row: dict[str, Any]) -> str:
    texto = _norm(f"{row.get('Nome', '')} {row.get('Descrição', '')}")
    if "femea" in texto or "fêmea" in texto or "soquete" in texto or "socket" in texto:
        return "Fêmea"
    if "macho" in texto or "plug" in texto or "pino" in texto:
        return "Macho"
    return "Fêmea"


def _extrair_material_revestimento(row: dict[str, Any]) -> str:
    texto = _norm(f"{row.get('Nome', '')} {row.get('Descrição', '')}")
    if any(k in texto for k in ("poliamida", "policarbonato", "plastico", "plástico", "pvc", "pa66")):
        return "Plástico"
    if any(k in texto for k in ("metal", "cobre", "latão", "latao", "aço", "aco")):
        return "Metal"
    return "Plástico"


def _norm(text: str) -> str:
    text = str(text or "").strip().lower()
    text = text.replace("\n", " ")
    return re.sub(r"\s+", " ", text)


def _map_headers(header_row: tuple[Any, ...]) -> dict[str, int]:
    """Mapa chave lógica → índice de coluna (0-based)."""
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        n = _norm(cell)
        for prefix, key in HEADER_ALIASES:
            if n.startswith(prefix) or n == prefix.rstrip(":"):
                mapping.setdefault(key, idx)
                break
        # match exato curto
        if n == "fotos":
            mapping.setdefault("fotos", idx)
        if n == "sku":
            mapping.setdefault("sku", idx)
        if n == "marca":
            mapping.setdefault("marca", idx)
        if n == "modelo":
            mapping.setdefault("modelo", idx)
    return mapping


def _find_header_row(ws) -> int:
    """Retorna índice 1-based da linha de headers (geralmente 3)."""
    for r in range(1, 8):
        vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column, 40) + 1)]
        joined = " ".join(_norm(v) for v in vals if v)
        if "sku" in joined and ("preço" in joined or "preco" in joined):
            return r
    return 3


def _find_data_start(ws, header_row: int, colmap: dict[str, int] | None = None) -> int:
    """
    Primeira linha de dados após ajuda.

    Pula linhas com MergedCell nas colunas críticas (SKU/Título) — no template
    ML a row 7 costuma ser faixa mesclada de UI.
    """
    start = header_row + 4  # após Obrigatório + ~3 linhas de ajuda
    probe_cols = []
    if colmap:
        for key in ("sku", "titulo", "preco", "marca"):
            if key in colmap:
                probe_cols.append(colmap[key] + 1)
    if not probe_cols:
        probe_cols = [1, 6, 8]

    for r in range(start, header_row + 12):
        # Se alguma coluna crítica é MergedCell, pula a linha
        if any(isinstance(ws.cell(r, c), MergedCell) for c in probe_cols):
            continue
        v = ws.cell(r, 1).value
        text = _norm(v)
        if text in {"obrigatório", "obrigatorio"}:
            continue
        if text and any(
            x in text
            for x in (
                "buscar",
                "máximo",
                "maximo",
                "este dado",
                "obrigat",
                "informe o produto",
                "digite apenas",
            )
        ):
            continue
        return r
    return start + 1


def _estoque(raw: Any) -> Any:
    if raw in (None, "", "nan"):
        return 2
    return _as_int_cm(raw)


def _as_int_cm(raw: Any, *, min_val: int | None = None) -> Any:
    if raw in (None, "", "nan"):
        return ""
    try:
        val = int(round(float(str(raw).replace(",", "."))))
    except (TypeError, ValueError):
        return ""
    if min_val is not None and val < min_val:
        return min_val
    return val


def _as_float(
    raw: Any, ndigits: int = 2, *, min_val: float | None = None
) -> Any:
    if raw in (None, "", "nan"):
        return ""
    try:
        val = round(float(str(raw).replace(",", ".")), ndigits)
    except (TypeError, ValueError):
        return ""
    if min_val is not None and val < min_val:
        return min_val
    return val


def _gtin_ml(row: dict[str, Any]) -> str:
    for campo in ("Código de barras", "GTIN", "EAN"):
        raw = str(row.get(campo) or "").strip()
        if raw.lower() in ("", "nan", "none"):
            continue
        m = re.search(r"\b(\d{8,14})\b", raw)
        if m:
            return m.group(1)
    return "O produto não tem código cadastrado"


def _medida_embalagem_cm(row: dict[str, Any], campo: str) -> int:
    val = _as_int_cm(row.get(campo, ""), min_val=1)
    if isinstance(val, int) and val >= 1:
        return val
    return 1


def _peso_embalagem_kg(row: dict[str, Any]) -> float:
    val = _as_float(row.get("Peso (kg)", ""), 3, min_val=0.001)
    if isinstance(val, (int, float)) and val > 0:
        return float(val)
    return 0.01


def _fotos_cell(raw: Any) -> str:
    if raw in (None, "", "nan"):
        return ""
    parts = [
        p.strip()
        for p in re.split(r"[\n,;]+", str(raw))
        if p.strip().startswith("http")
    ]
    return ",".join(parts)


def _titulo(nome: str) -> str:
    t = re.sub(r"\s+", " ", str(nome or "")).strip()
    if len(t) <= 60:
        return t
    return t[:60].rsplit(" ", 1)[0].strip()


def _produto_para_valores(row: dict[str, Any], *, sheet: str = "Outros") -> dict[str, Any]:
    titulo = _titulo(row.get("Nome", ""))
    sku = str(row.get("SKU", "")).strip()
    valores: dict[str, Any] = {
        "titulo": titulo,
        "qtd_chars": len(titulo),
        "condicao": "Novo",
        "gtin": _gtin_ml(row),
        "varia_cor": _extrair_cor_ml(row),
        "fotos": _fotos_cell(row.get("Fotos", "")),
        "sku": sku,
        "estoque": _estoque(row.get("Estoque", "")),
        "preco": _as_float(row.get("Preço", "")),
        "formato_venda": "Unidade",
        "qtd_unidades": 1,
        "descricao": str(row.get("Descrição", "") or ""),
        "largura": _medida_embalagem_cm(row, "Largura (cm)"),
        "altura": _medida_embalagem_cm(row, "Altura (cm)"),
        "profundidade": _medida_embalagem_cm(row, "Comprimento (cm)"),
        "peso": _peso_embalagem_kg(row),
        "tipo_anuncio": "Clássico",
        "forma_envio": "Mercado Envios",
        "custo_envio": "",
        "retirar": "Não",
        "tipo_garantia": "Garantia do vendedor",
        "tempo_garantia": 12,
        "unidade_garantia": "meses",
        "marca": "Wago",
        "fabricante": "Wago",
        "modelo": sku,
        "modelo_alfanumerico": sku,
    }
    voltagem = _extrair_voltagem_ml(row)
    if voltagem:
        valores["varia_voltagem"] = voltagem
    if sheet == "Conectores":
        valores["tipo_produto"] = "Conector"
        valores["genero_conector"] = _extrair_genero_conector(row)
        valores["material_revestimento"] = _extrair_material_revestimento(row)
    elif sheet == "Relés":
        valores["tipo_produto"] = "Relé"
    elif sheet == "Contatores":
        valores["tipo_produto"] = "Contator"
    elif sheet == "Interruptores Automáticos":
        valores["tipo_produto"] = "Automático"
    elif sheet == "Fichas Elétricas":
        valores["tipo_conexao_ficha"] = _extrair_genero_conector(row)
    return valores


# Colunas com fórmula ou dropdown já preenchido pelo template.
# qtd_chars, tarifa e BUYBOX são fórmulas. Retirar e custo de envio
# permanecem no valor oficial da planilha ("Não aceito", frete do template).
_NAO_GRAVAR = {"qtd_chars", "custo_envio", "retirar"}

_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _qn(tag: str) -> str:
    return f"{{{_NS}}}{tag}"


def _col_letter(index: int) -> str:
    """Índice 1-based para a letra da coluna."""
    letters = ""
    while index:
        index, resto = divmod(index - 1, 26)
        letters = chr(65 + resto) + letters
    return letters


def _col_index(letters: str) -> int:
    total = 0
    for char in letters:
        total = total * 26 + (ord(char.upper()) - 64)
    return total


def _caminho_da_aba(xlsx: Path, aba: str) -> str:
    with zipfile.ZipFile(xlsx) as arquivo:
        livro = etree.fromstring(arquivo.read("xl/workbook.xml"))
        rels = etree.fromstring(arquivo.read("xl/_rels/workbook.xml.rels"))
    rid = None
    for sheet in livro.findall(f"{_qn('sheets')}/{_qn('sheet')}"):
        if sheet.get("name") == aba:
            rid = sheet.get(f"{{{_NS_R}}}id")
            break
    if not rid:
        raise SystemExit(f"Aba {aba!r} não existe no workbook.")
    for rel in rels:
        if rel.get("Id") == rid:
            alvo = rel.get("Target") or ""
            if alvo.startswith("/"):
                return alvo.lstrip("/")
            if not alvo.startswith("xl/"):
                return "xl/" + alvo.lstrip("/")
            return alvo
    raise SystemExit(f"Aba {aba!r} sem arquivo XML.")


def _linhas_para_produtos(
    ws, data_start: int, sku_col: int, quantidade: int
) -> tuple[list[int], set[str]]:
    livres: list[int] = []
    ocupados: set[str] = set()
    linha = data_start
    while len(livres) < quantidade and linha < data_start + 5000:
        valor = ws.cell(linha, sku_col).value
        if valor is None or str(valor).strip() == "":
            livres.append(linha)
        else:
            ocupados.add(str(valor).strip().upper())
        linha += 1
    return livres, ocupados


def _formulas(ws, linha: int, colmap: dict[str, int]) -> set[str]:
    refs: set[str] = set()
    for key, col in colmap.items():
        if key == "fotos":
            continue
        cell = ws.cell(linha, col + 1)
        if cell.data_type == "f" or (
            isinstance(cell.value, str) and cell.value.startswith("=")
        ):
            refs.add(cell.coordinate)
    return refs


class _SharedStrings:
    """Tabela de strings compartilhadas (`xl/sharedStrings.xml`)."""

    def __init__(self, xml: bytes) -> None:
        self.root = etree.fromstring(xml)
        self._indice: dict[str, int] = {}
        for idx, si in enumerate(self.root.findall(_qn("si"))):
            t = si.find(_qn("t"))
            if t is not None and len(si) == 1:
                self._indice.setdefault(t.text or "", idx)
        try:
            self._count = int(self.root.get("count") or len(self._indice))
        except ValueError:
            self._count = len(self._indice)

    def indice(self, texto: str) -> int:
        self._count += 1
        existente = self._indice.get(texto)
        if existente is not None:
            return existente
        novo = len(self.root.findall(_qn("si")))
        si = etree.SubElement(self.root, _qn("si"))
        t = etree.SubElement(si, _qn("t"))
        if texto[:1].isspace() or texto[-1:].isspace() or "\n" in texto:
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = texto
        self._indice[texto] = novo
        return novo

    def to_bytes(self) -> bytes:
        total_unicos = len(self.root.findall(_qn("si")))
        self.root.set("count", str(max(self._count, total_unicos)))
        self.root.set("uniqueCount", str(total_unicos))
        corpo = etree.tostring(
            self.root, xml_declaration=False, encoding="UTF-8"
        )
        return b'<?xml version="1.0" encoding="UTF-8"?>\n' + corpo


def _estilos_colunas(root: etree._Element) -> dict[int, str]:
    estilos: dict[int, str] = {}
    cols = root.find(_qn("cols"))
    if cols is None:
        return estilos
    for col in cols.findall(_qn("col")):
        estilo = col.get("style")
        if not estilo:
            continue
        try:
            inicio = int(col.get("min") or 0)
            fim = int(col.get("max") or 0)
        except ValueError:
            continue
        for indice in range(inicio, fim + 1):
            estilos[indice] = estilo
    return estilos


def _aplicar_celulas(
    xml: bytes,
    updates: dict[str, Any],
    *,
    shared_strings: _SharedStrings | None = None,
    forcar_Refs: set[str] | None = None,
) -> bytes:
    etree.register_namespace("r", _NS_R)
    root = etree.fromstring(xml)
    dados = root.find(_qn("sheetData"))
    if dados is None:
        raise SystemExit("A aba não tem sheetData.")
    estilos_col = _estilos_colunas(root)
    por_linha: dict[str, list[tuple[str, Any]]] = {}
    for ref, valor in updates.items():
        linha = "".join(ch for ch in ref if ch.isdigit())
        por_linha.setdefault(linha, []).append((ref, valor))
    linhas = {row.get("r"): row for row in dados.findall(_qn("row"))}
    for linha, pares in por_linha.items():
        row = linhas.get(linha)
        if row is None:
            row = etree.Element(_qn("row"))
            row.set("r", linha)
            dados.append(row)
            linhas[linha] = row
        for ref, valor in pares:
            _upsert_celula(
                row,
                ref,
                valor,
                estilos_col=estilos_col,
                shared_strings=shared_strings,
                forcar_formula=bool(forcar_Refs and ref in forcar_Refs),
            )
    corpo = etree.tostring(root, xml_declaration=False, encoding="UTF-8")
    return b'<?xml version="1.0" encoding="UTF-8"?>\n' + corpo


def _upsert_celula(
    row: etree._Element,
    ref: str,
    valor: Any,
    *,
    estilos_col: dict[int, str] | None = None,
    shared_strings: _SharedStrings | None = None,
    forcar_formula: bool = False,
) -> None:
    celula = None
    for candidata in row.findall(_qn("c")):
        if candidata.get("r") == ref:
            celula = candidata
            break
    if (
        celula is not None
        and celula.find(_qn("f")) is not None
        and not forcar_formula
    ):
        return
    indice = _col_index("".join(ch for ch in ref if ch.isalpha()))
    if celula is None:
        celula = etree.Element(_qn("c"))
        celula.set("r", ref)
        inserida = False
        for candidata in row.findall(_qn("c")):
            outra = "".join(ch for ch in (candidata.get("r") or "") if ch.isalpha())
            if _col_index(outra) > indice:
                candidata.addprevious(celula)
                inserida = True
                break
        if not inserida:
            row.append(celula)
    if not celula.get("s"):
        estilo = (estilos_col or {}).get(indice, "0")
        celula.set("s", estilo)
    for filho in list(celula):
        if filho.tag in {_qn("v"), _qn("is"), _qn("f")}:
            celula.remove(filho)
    if isinstance(valor, bool):
        valor = str(valor)
    if isinstance(valor, (int, float)):
        celula.set("t", "n")
        numero = etree.SubElement(celula, _qn("v"))
        numero.text = str(int(valor)) if isinstance(valor, int) else str(valor)
        return
    conteudo = "" if valor is None else str(valor)
    if shared_strings is not None:
        celula.set("t", "s")
        idx_str = shared_strings.indice(conteudo)
        v = etree.SubElement(celula, _qn("v"))
        v.text = str(idx_str)
        return
    celula.set("t", "inlineStr")
    bloco = etree.SubElement(celula, _qn("is"))
    texto = etree.SubElement(bloco, _qn("t"))
    if conteudo[:1].isspace() or conteudo[-1:].isspace():
        texto.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    texto.text = conteudo


def _substituir_partes(
    origem: Path, destino: Path, partes: dict[str, bytes]
) -> None:
    with zipfile.ZipFile(origem, "r") as entrada, zipfile.ZipFile(
        destino, "w"
    ) as saida:
        for info in entrada.infolist():
            dados = partes.get(info.filename)
            if dados is None:
                dados = entrada.read(info.filename)
            novo = zipfile.ZipInfo(filename=info.filename, date_time=info.date_time)
            novo.compress_type = info.compress_type
            novo.external_attr = info.external_attr
            saida.writestr(novo, dados)


def _contagem(xlsx: Path, prefixo: str) -> int:
    with zipfile.ZipFile(xlsx) as arquivo:
        return sum(1 for nome in arquivo.namelist() if nome.startswith(prefixo))


def _validar_saida(
    caminho: Path,
    *,
    template: Path,
    planos: dict[str, dict[str, Any]],
) -> None:
    if _contagem(caminho, "xl/media/") != _contagem(template, "xl/media/"):
        raise SystemExit("A gravação removeu imagens do workbook.")
    with zipfile.ZipFile(caminho) as arquivo:
        for aba, info in planos.items():
            xml = arquivo.read(info["parte"])
            if b"<dataValidations" not in xml:
                raise SystemExit(f"A aba {aba} perdeu as validações.")
            if b"<f>" not in xml and b"<f " not in xml:
                raise SystemExit(f"A aba {aba} perdeu fórmulas.")
            if b't="inlineStr"' in xml:
                raise SystemExit(
                    f"A aba {aba} contém inlineStr em vez de sharedStrings."
                )
    wb = load_workbook(caminho, read_only=True, data_only=False)
    try:
        for aba, info in planos.items():
            if aba not in wb.sheetnames:
                raise SystemExit(f"Aba {aba!r} sumiu do arquivo gerado.")
            ws = wb[aba]
            vistos: set[str] = set()
            for row in ws.iter_rows(min_row=1, max_row=120, values_only=True):
                for celula in row:
                    if celula is not None:
                        vistos.add(str(celula).strip())
            faltando = [sku for sku in info["skus"] if sku not in vistos]
            if faltando:
                raise SystemExit(
                    f"SKU ausente na aba {aba}: {', '.join(faltando)}"
                )
    finally:
        wb.close()


def _imprimir_rejeicoes(rejeitados: list[tuple[str, str]]) -> None:
    for sku, motivo in rejeitados:
        print(f"  - {sku}: {motivo}")


def _publicar_resumo_ml(dest: Path, resumo: dict[str, Any]) -> None:
    path = dest.with_suffix(".resumo.json")
    tmp = path.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(resumo, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)


def _eh_automatico(sheet: str | None) -> bool:
    if sheet is None:
        return True
    s = _norm(sheet)
    return s in {"", "auto", "automatico", "automático"}


def fill_ml(
    master_csv: Path | str,
    *,
    sheet: str | None = ABA_AUTOMATICA,
    output: Path | str | None = None,
    template: Path | str | None = None,
    dry_run: bool = False,
    force: bool = False,
) -> Path | None:
    """
    Preenche o Anunciar-ML.xlsx sem regravar o workbook inteiro.

    Quando ``sheet`` é ``"Automático"`` (ou vazio/``None``), cada produto é
    direcionado para a aba correspondente à sua categoria/título. Quando uma
    aba específica é informada, todos os produtos vão para aquela aba.
    """
    ensure_dirs()
    tpl = Path(template) if template else TEMPLATE
    if not tpl.exists():
        raise SystemExit(f"Template ML ausente: {tpl}")

    master_path = Path(master_csv)
    if not master_path.exists():
        raise SystemExit(f"Master não encontrado: {master_path}")

    df = pd.read_csv(master_path, sep=";", dtype=str, encoding="utf-8-sig").fillna("")
    if df.empty:
        raise SystemExit("Master sem produtos.")

    aceitos, rejeitados = selecionar_para_exportacao(df.to_dict(orient="records"))
    if output:
        out_path = Path(output)
    else:
        stamp = datetime.now().strftime("%d-%m_%H%M%S")
        out_path = OUT / f"Anunciar-ML-preenchido_{stamp}.xlsx"
    if out_path.resolve() == tpl.resolve():
        raise SystemExit("A saída não pode ser o template oficial.")
    if out_path.exists() and not force and not dry_run:
        raise SystemExit(
            f"Arquivo já existe: {out_path}. Use --force para substituir."
        )

    modo_auto = _eh_automatico(sheet)
    origem = out_path if out_path.exists() else tpl
    wb = load_workbook(origem)
    try:
        abas_validas = [
            n
            for n in wb.sheetnames
            if n not in {"Ajuda", "extra info", "Textos jurídicos"}
        ]
        if not modo_auto and sheet not in wb.sheetnames:
            raise SystemExit(
                f"Aba {sheet!r} não existe. Disponíveis: {wb.sheetnames}"
            )

        por_aba: dict[str, list[dict[str, Any]]] = {}
        for row in aceitos:
            aba_alvo = (
                inferir_aba_ml(row, abas_validas) if modo_auto else str(sheet)
            )
            por_aba.setdefault(aba_alvo, []).append(row)

        planos: dict[str, dict[str, Any]] = {}
        total_gravar = 0
        for aba_alvo, itens_aba in por_aba.items():
            ws = wb[aba_alvo]
            header_row = _find_header_row(ws)
            headers = tuple(
                ws.cell(header_row, c).value
                for c in range(1, min(ws.max_column, 56) + 1)
            )
            colmap = _map_headers(headers)
            if "titulo" not in colmap or "sku" not in colmap:
                raise SystemExit(
                    f"Não foi possível mapear colunas na aba {aba_alvo!r} "
                    f"(header row {header_row})."
                )
            data_start = _find_data_start(ws, header_row, colmap)
            livres, ocupados = _linhas_para_produtos(
                ws, data_start, colmap["sku"] + 1, len(itens_aba)
            )
            gravar_aba: list[dict[str, Any]] = []
            for row in itens_aba:
                sku = str(row["SKU"]).strip().upper()
                if sku in ocupados:
                    rejeitados.append((sku, f"SKU já presente na aba {aba_alvo}"))
                    continue
                gravar_aba.append(row)
                ocupados.add(sku)
            if not gravar_aba:
                continue
            if len(livres) < len(gravar_aba):
                raise SystemExit(
                    f"A aba {aba_alvo!r} não tem linhas livres suficientes."
                )

            updates: dict[str, Any] = {}
            forcar_refs: set[str] = set()
            for excel_row, row in zip(livres, gravar_aba):
                valores = _produto_para_valores(row, sheet=aba_alvo)
                formulas = _formulas(ws, excel_row, colmap)
                for key, col_idx in colmap.items():
                    if key in _NAO_GRAVAR or key not in valores:
                        continue
                    valor = valores[key]
                    if valor in ("", None):
                        continue
                    ref = f"{_col_letter(col_idx + 1)}{excel_row}"
                    if ref in formulas:
                        continue
                    if key == "fotos":
                        forcar_refs.add(ref)
                    updates[ref] = valor
            parte = _caminho_da_aba(origem, aba_alvo)
            planos[aba_alvo] = {
                "parte": parte,
                "updates": updates,
                "forcar_refs": forcar_refs,
                "skus": [str(row["SKU"]).strip().upper() for row in gravar_aba],
                "header_row": header_row,
                "data_start": data_start,
                "colmap": colmap,
                "exportados": len(gravar_aba),
            }
            total_gravar += len(gravar_aba)
    finally:
        wb.close()

    abas_resumo = {aba: info["exportados"] for aba, info in planos.items()}
    aba_label = (
        ", ".join(f"{k}({v})" for k, v in abas_resumo.items())
        if modo_auto and abas_resumo
        else (str(sheet) if not modo_auto else ABA_AUTOMATICA)
    )
    resumo = {
        "origem": str(master_path),
        "aba": str(sheet or ABA_AUTOMATICA),
        "abas_preenchidas": abas_resumo,
        "lidos": int(len(df)),
        "exportados": total_gravar,
        "rejeitados": len(rejeitados),
        "motivos": [{"sku": sku, "motivo": motivo} for sku, motivo in rejeitados],
    }
    print(
        f"ML: lidos={resumo['lidos']} exportados={resumo['exportados']} "
        f"rejeitados={resumo['rejeitados']} aba={aba_label}"
    )
    _imprimir_rejeicoes(rejeitados)
    if total_gravar == 0:
        raise SystemExit("Nenhum produto exportado para o Mercado Livre.")
    if dry_run:
        print("Simulação: nenhum arquivo foi gravado.")
        return None

    out_path.parent.mkdir(parents=True, exist_ok=True)
    temporario = out_path.with_name(out_path.stem + ".parcial.xlsx")
    try:
        with zipfile.ZipFile(origem) as arquivo:
            ss_xml = (
                arquivo.read("xl/sharedStrings.xml")
                if "xl/sharedStrings.xml" in arquivo.namelist()
                else None
            )
            shared_strings = (
                _SharedStrings(ss_xml) if ss_xml is not None else None
            )
            partes_novas: dict[str, bytes] = {}
            for aba_alvo, info in planos.items():
                xml_aba = arquivo.read(info["parte"])
                partes_novas[info["parte"]] = _aplicar_celulas(
                    xml_aba,
                    info["updates"],
                    shared_strings=shared_strings,
                    forcar_Refs=info["forcar_refs"],
                )
        if shared_strings is not None:
            partes_novas["xl/sharedStrings.xml"] = shared_strings.to_bytes()
        _substituir_partes(origem, temporario, partes_novas)
        _validar_saida(
            temporario,
            template=tpl,
            planos=planos,
        )
        temporario.replace(out_path)
    except BaseException:
        temporario.unlink(missing_ok=True)
        raise
    _publicar_resumo_ml(out_path, resumo)
    print(f"ML preenchido: {out_path}")
    for aba_alvo, info in planos.items():
        print(
            f"  Aba: {aba_alvo} | Header row: {info['header_row']} | "
            f"Dados a partir da row {info['data_start']} | "
            f"Produtos: {info['exportados']}"
        )

    if output is None:
        latest = OUT / (
            "Anunciar-ML-preenchido.xlsx"
            if modo_auto or sheet == "Outros"
            else f"Anunciar-ML-preenchido-{sheet}.xlsx"
        )
        if latest.resolve() != out_path.resolve():
            latest.write_bytes(out_path.read_bytes())
            print(f"Cópia latest: {latest}")
    return out_path
