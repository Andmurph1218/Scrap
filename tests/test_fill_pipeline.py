"""Testes das etapas 4a / 4b / 5 (fill templates + pastas)."""

from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from wago.fill_ml import _map_headers, _produto_para_valores, fill_ml  # noqa: E402
from wago.fill_nuvemshop import COLUNAS, fill_nuvemshop  # noqa: E402
from wago.images import pasta_produto  # noqa: E402


def _mini_master(tmp: Path) -> Path:
    path = tmp / "master.csv"
    df = pd.DataFrame(
        [
            {
                "Identificador URL": "conector-221-413-wago",
                "Nome": "Conector de emenda 221-413 Wago",
                "Categorias": "Bornes > Conectores",
                "Preço": "13.13",
                "Peso (kg)": "0.01",
                "Altura (cm)": "2",
                "Largura (cm)": "3",
                "Comprimento (cm)": "4",
                "Estoque": "2",
                "SKU": "221-413",
                "Descrição": "Produto teste",
                "Tags": "221-413, Wago",
                "Título para SEO": "Conector de emenda 221-413 Wago",
                "Descrição para SEO": "Produto teste",
                "Marca": "Wago",
                "Produto Físico": "SIM",
                "MPN (Cód. Exclusivo Modelo Fabricante)": "221-413",
                "Custo": "7.18",
                "Visibilidade": "Visível",
                "Exibir na loja": "SIM",
                "Frete gratis": "NÃO",
                "Fotos": "https://www.wago.com/medias/example.jpg",
                "IPI": "0.7",
                "NCM": "",
            }
        ]
    )
    # completa colunas nuvemshop vazias
    for c in COLUNAS:
        if c not in df.columns:
            df[c] = ""
    df.to_csv(path, sep=";", index=False, encoding="utf-8-sig")
    return path


def test_pasta_produto():
    assert pasta_produto("221-413", "Conector de emenda 221-413 Wago").startswith(
        "221-413__"
    )


def test_fill_nuvemshop(tmp_path, monkeypatch):
    monkeypatch.setattr("wago.export.margin_divisor", lambda: 0.6)
    master = _mini_master(tmp_path)
    out = tmp_path / "ns.csv"
    fill_nuvemshop(master, output=out)
    df = pd.read_csv(out, sep=";", dtype=str)
    assert list(df.columns) == COLUNAS
    assert len(df) == 1
    assert df.iloc[0]["SKU"] == "221-413"
    assert df.iloc[0]["Marca"] == "Wago"
    # sem linhas de exemplo do template
    assert "Shorts" not in df["Nome"].tolist()


def test_map_headers_outros():
    headers = (
        "Título: informe o produto, marca, modelo",
        "Quantidade de caracteres",
        "Condição",
        "Código universal de produto",
        "Fotos",
        "SKU",
        "Estoque no depósito",
        "Preço [R$]",
        "Descrição",
        "Largura (cm)",
        "Altura (cm)",
        "Profundidade (cm)",
        "Peso físico (kg) \nEmbalagem",
        "Tipo de anúncio",
        "Marca",
        "Modelo",
    )
    m = _map_headers(headers)
    assert m["titulo"] == 0
    assert m["sku"] == 5
    assert m["preco"] == 7
    assert m["marca"] == 14


def test_produto_valores_titulo_60():
    vals = _produto_para_valores(
        {
            "Nome": "A" * 80 + " Wago",
            "SKU": "X",
            "Preço": "10",
            "Estoque": "2",
            "Fotos": "https://a.com/1.jpg, https://a.com/2.jpg",
        }
    )
    assert len(vals["titulo"]) <= 60
    assert vals["condicao"] == "Novo"
    assert vals["marca"] == "Wago"
    assert "https://a.com/1.jpg" in vals["fotos"]


def test_fill_ml_outros(tmp_path, monkeypatch):
    monkeypatch.setattr("wago.export.margin_divisor", lambda: 0.6)
    master = _mini_master(tmp_path)
    out = tmp_path / "ml.xlsx"
    fill_ml(master, sheet="Outros", output=out)
    wb = load_workbook(out, read_only=True, data_only=True)
    assert "Outros" in wb.sheetnames
    ws = wb["Outros"]
    # procura SKU na planilha
    found = False
    for row in ws.iter_rows(min_row=1, max_row=40, values_only=True):
        if row and "221-413" in [str(c) for c in row if c is not None]:
            found = True
            break
    wb.close()
    assert found, "SKU 221-413 não encontrado na aba Outros"
