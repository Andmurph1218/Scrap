"""Fase 4: fluxo sintético, reabertura das saídas e ciclo do painel.

Não usa login, sitemap nem scraping ao vivo.
"""

from __future__ import annotations

import json
import sys
import threading
import urllib.error
import urllib.request
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pandas as pd
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from wago.dashboard.server import criar_servidor, servir  # noqa: E402
from wago.dashboard.service import Painel  # noqa: E402
from wago.dashboard.store import Store  # noqa: E402
from wago.export import (  # noqa: E402
    carregar_checkpoint_master,
    gravar_checkpoint_master,
)
from wago.fill_ml import fill_ml  # noqa: E402
from wago.fill_nuvemshop import COLUNAS, fill_nuvemshop  # noqa: E402
from wago.images import baixar_imagens_lote, pasta_produto  # noqa: E402
from wago.scrape import recuperar_checkpoint  # noqa: E402

TEMPLATE_NS = ROOT / "templates" / "produtos-Nuvemshop.csv"
TEMPLATE_ML = ROOT / "templates" / "Anunciar-ML.xlsx"
JPEG = b"\xff\xd8\xff" + b"\x00" * 32
_NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_NS_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"


def _isolado(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr("wago.export.STATE", tmp_path)
    monkeypatch.setattr("wago.export.margin_divisor", lambda: 0.6)
    monkeypatch.setattr("wago.state.SKUS_PROCESSADOS", tmp_path / "skus_processados.txt")


def _produtos() -> list[dict]:
    return [
        {
            "SKU": " 221-413 ",
            "Nome": "Conector de Emenda Ção",
            "Categorias": "Bornes > Conectores",
            "Preço": "13.13",
            "Peso (kg)": "0.01",
            "Altura (cm)": "2",
            "Largura (cm)": "3",
            "Comprimento (cm)": "4",
            "Estoque": "2",
            "Descrição": "Produto sintético",
            "Custo": "7.18",
            "IPI": "0.70",
            "Fotos": "https://cdn.example/a.png",
        },
        {
            "SKU": "222-412",
            "Nome": "Borne laranja",
            "Categorias": "Bornes",
            "Preço": "16.67",
            "Peso (kg)": "0.02",
            "Altura (cm)": "1",
            "Largura (cm)": "1",
            "Comprimento (cm)": "2",
            "Estoque": "2",
            "Descrição": "IPI zero",
            "Custo": "10",
            "IPI": "0",
            "Fotos": "https://cdn.example/b.png",
        },
        {
            "SKU": "sem-ipi",
            "Nome": "Sem imposto",
            "Preço": "10.00",
            "Custo": "6",
            "IPI": "",
            "Peso (kg)": "0.01",
            "Altura (cm)": "1",
            "Largura (cm)": "1",
            "Comprimento (cm)": "1",
        },
    ]


def _master_do_checkpoint(tmp_path, monkeypatch) -> Path:
    _isolado(tmp_path, monkeypatch)
    gravar_checkpoint_master(_produtos(), 7)
    vistos: set[str] = set()
    recuperados = recuperar_checkpoint(7, vistos)
    assert [item["SKU"] for item in recuperados] == ["221-413", "222-412"]
    assert "SEM-IPI" not in vistos
    reaberto = carregar_checkpoint_master(7)
    assert {item["SKU"] for item in reaberto} == {"221-413", "222-412", "SEM-IPI"}
    destino = tmp_path / "master-reaberto.csv"
    pd.DataFrame(reaberto).to_csv(destino, sep=";", index=False, encoding="utf-8-sig")
    return destino


def _xml_da_aba(caminho: Path, aba: str) -> bytes:
    with zipfile.ZipFile(caminho) as arquivo:
        livro = ET.fromstring(arquivo.read("xl/workbook.xml"))
        rels = ET.fromstring(arquivo.read("xl/_rels/workbook.xml.rels"))
        rid = ""
        for sheet in livro.findall(f"{{{_NS_MAIN}}}sheets/{{{_NS_MAIN}}}sheet"):
            if sheet.attrib.get("name") == aba:
                rid = sheet.attrib.get(f"{{{_NS_REL}}}id", "")
        alvo = ""
        for rel in rels.findall(f"{{{_NS_PKG}}}Relationship"):
            if rel.attrib.get("Id") == rid:
                alvo = "xl/" + rel.attrib["Target"].lstrip("/")
        if not alvo:
            raise AssertionError(f"Aba {aba} ausente em {caminho.name}")
        return arquivo.read(alvo)


def _contagem(caminho: Path, prefixo: str) -> int:
    with zipfile.ZipFile(caminho) as arquivo:
        return sum(1 for nome in arquivo.namelist() if nome.startswith(prefixo))


class _Sessao:
    def __init__(self, conteudo: bytes):
        self.conteudo = conteudo
        self.chamadas = 0

    def get(self, url, timeout=30):
        self.chamadas += 1
        return type("Resposta", (), {"status_code": 200, "content": self.conteudo})()


class _Recusa:
    def get(self, url, timeout=30):
        raise AssertionError(f"download repetido: {url}")


def test_saidas_reabertas_depois_da_geracao(tmp_path, monkeypatch):
    master = _master_do_checkpoint(tmp_path, monkeypatch)
    csv_path = tmp_path / "produtos-Nuvemshop-preenchido.csv"
    fill_nuvemshop(master, output=csv_path)

    cabecalho = pd.read_csv(TEMPLATE_NS, sep=";", nrows=0, encoding="utf-8-sig").columns.tolist()
    reaberto = pd.read_csv(csv_path, sep=";", dtype=str, encoding="utf-8-sig").fillna("")
    assert list(reaberto.columns) == cabecalho
    assert len(reaberto.columns) == 31
    assert set(reaberto.columns) == set(COLUNAS)
    assert reaberto["SKU"].tolist() == ["221-413", "222-412"]
    linha = reaberto.loc[reaberto["SKU"] == "221-413"].iloc[0]
    assert linha["Identificador URL"] == "conector-de-emenda-cao"
    assert linha["Preço"] == "13.13"
    assert linha["Custo"] == "7.18"
    assert linha["Marca"] == "Wago"
    assert linha["Exibir na loja"] == "SIM"
    zero = reaberto.loc[reaberto["SKU"] == "222-412"].iloc[0]
    assert zero["Preço"] == "16.67"
    assert zero["Custo"] == "10.00"
    assert zero["Identificador URL"] == "borne-laranja"
    resumo = json.loads(csv_path.with_suffix(".resumo.json").read_text(encoding="utf-8"))
    assert resumo["exportados"] == 2
    assert any(item["sku"] == "SEM-IPI" for item in resumo["motivos"])

    imagens = tmp_path / "imagens"
    sessao = _Sessao(JPEG)
    baixar_imagens_lote(master, output_dir=imagens, session=sessao, tentativas=1)
    pasta = imagens / pasta_produto("221-413", "Conector de Emenda Ção")
    arquivo = pasta / "01.jpg"
    assert arquivo.read_bytes().startswith(b"\xff\xd8\xff")
    assert sessao.chamadas == 2
    baixar_imagens_lote(master, output_dir=imagens, session=_Recusa(), tentativas=1)
    assert arquivo.read_bytes().startswith(b"\xff\xd8\xff")
    relatorio = json.loads((imagens / "resumo-imagens.json").read_text(encoding="utf-8"))
    assert relatorio["baixadas"] == 0
    assert relatorio["ignoradas"] == 2

    outros = tmp_path / "Anunciar-ML-Outros.xlsx"
    conectores = tmp_path / "Anunciar-ML-Conectores.xlsx"
    fill_ml(master, sheet="Outros", output=outros)
    fill_ml(master, sheet="Conectores", output=conectores)
    xml_outros = _xml_da_aba(outros, "Outros")
    xml_conectores = _xml_da_aba(conectores, "Conectores")
    xml_outros_intocado = _xml_da_aba(outros, "Conectores")
    with zipfile.ZipFile(outros) as z_out, zipfile.ZipFile(conectores) as z_con:
        ss_outros = z_out.read("xl/sharedStrings.xml")
        ss_conectores = z_con.read("xl/sharedStrings.xml")
    assert b't="inlineStr"' not in xml_outros
    assert b't="inlineStr"' not in xml_conectores
    assert b'r="F8" s="0" t="s"' in xml_outros
    assert b'r="F9" s="0" t="s"' in xml_outros
    assert b'r="H9" s="0" t="s"' in xml_conectores
    assert b'r="H9"' not in xml_outros_intocado
    assert b"221-413" in ss_outros
    assert b"222-412" in ss_outros
    assert b"222-412" in ss_conectores
    assert b"<f>" in xml_outros or b"<f " in xml_outros
    assert b"<dataValidations" in xml_outros
    assert _contagem(outros, "xl/media/") == _contagem(TEMPLATE_ML, "xl/media/")

    livro = load_workbook(outros, read_only=True, data_only=False)
    try:
        aba = livro["Outros"]
        textos = []
        formulas = 0
        for row in aba.iter_rows(min_row=1, max_row=20):
            for celula in row:
                if isinstance(celula.value, str) and celula.value.startswith("="):
                    formulas += 1
                if celula.value is not None:
                    textos.append(str(celula.value))
        assert formulas > 0
        assert "221-413" in textos
        assert "222-412" in textos
    finally:
        livro.close()


def _post(porta: int, caminho: str, corpo: dict) -> tuple[int, dict]:
    pedido = urllib.request.Request(
        f"http://127.0.0.1:{porta}{caminho}",
        data=json.dumps(corpo).encode(),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(pedido) as resposta:
            return resposta.status, json.loads(resposta.read().decode())
    except urllib.error.HTTPError as erro:
        return erro.code, json.loads(erro.read().decode())


def test_dashboard_inicia_valida_e_encerra(tmp_path, monkeypatch):
    monkeypatch.setattr("wago.dashboard.service.OUT", tmp_path)
    monkeypatch.setattr("wago.dashboard.server.STATE", tmp_path / "estado")
    (tmp_path / "nota.csv").write_text("sku\n1\n", encoding="utf-8")
    with pytest.raises(SystemExit, match="Porta inválida"):
        servir(0)

    entrou = threading.Event()

    def lenta(params, cancelado):
        entrou.set()
        while not cancelado():
            threading.Event().wait(0.01)
        print("etapa encerrada apos cancelamento")
        return 2

    banco = tmp_path / "estado" / "dashboard.sqlite"
    painel = Painel(Store(banco), {"discover": lenta})
    httpd = criar_servidor(painel, 0)
    assert httpd.server_address[0] == "127.0.0.1"
    fio = threading.Thread(target=httpd.serve_forever, daemon=True)
    fio.start()
    porta = httpd.server_address[1]
    try:
        with pytest.raises(SystemExit, match="ocupada"):
            servir(porta)
        with urllib.request.urlopen(f"http://127.0.0.1:{porta}/") as resposta:
            assert b"Wago Pipeline" in resposta.read()
        codigo, dados = _post(porta, "/api/execucoes", {"operacao": "discover"})
        assert codigo == 400
        codigo, dados = _post(
            porta,
            "/api/execucoes",
            {"operacao": "fill-ml", "confirmar": True, "aba": "Aba Inexistente"},
        )
        assert codigo == 400
        with pytest.raises(urllib.error.HTTPError) as erro:
            urllib.request.urlopen(
                f"http://127.0.0.1:{porta}/api/arquivos/baixar?nome=../nota.csv"
            )
        assert erro.value.code == 400
        with urllib.request.urlopen(
            f"http://127.0.0.1:{porta}/api/arquivos/baixar?nome=nota.csv"
        ) as resposta:
            assert b"sku" in resposta.read()

        codigo, primeiro = _post(
            porta, "/api/execucoes", {"operacao": "discover", "confirmar": True}
        )
        assert codigo == 202
        assert entrou.wait(2)
        codigo, conflito = _post(
            porta, "/api/execucoes", {"operacao": "discover", "confirmar": True}
        )
        assert codigo == 409
        assert "andamento" in conflito["erro"]
        with urllib.request.urlopen(f"http://127.0.0.1:{porta}/api/painel") as resposta:
            assert json.loads(resposta.read().decode())["aplicacao"] == "executando"
        codigo, _ = _post(porta, f"/api/execucoes/{primeiro['id']}/cancelar", {})
        assert codigo == 200
        assert painel._thread is not None
        painel._thread.join(3)
    finally:
        httpd.shutdown()
        httpd.server_close()
        fio.join(3)

    with pytest.raises(urllib.error.URLError):
        urllib.request.urlopen(f"http://127.0.0.1:{porta}/", timeout=2)
    historico = Store(banco).obter_execucao(primeiro["id"])
    assert historico is not None
    assert historico["status"] == "cancelada"
    assert historico["codigo"] == 2
